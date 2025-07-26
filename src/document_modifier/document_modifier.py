"""
Document Modifier Module for AutoJobAgent.

This module provides the DocumentModifier class for generating and customizing
resumes and cover letters for job applications. It supports both .docx templates
and JSON-based content generation, with a fallback mechanism for PDF conversion.
"""

import json
import logging
from pathlib import Path
from typing import Dict, Any, Optional, List
from datetime import datetime

from docx import Document as DocxDocument
from docx2pdf import convert as docx2pdf_convert
from reportlab.lib.pagesizes import LETTER
from reportlab.pdfgen import canvas

from src.core.job_database import ModernJobDatabase
# Removed circular import: from src.utils.document_generator import customize

# --- Initialization ---
logger = logging.getLogger(__name__)


class DocumentModifier:
    """
    Generates and customizes documents (resume, cover letter) for job applications.
    Designed to be used as a worker in the job application pipeline.
    """

    def __init__(self, profile_name: str):
        """
        Initialize the DocumentModifier.

        Args:
            profile_name: The name of the profile to use for document generation.
        """
        if not profile_name:
            raise ValueError("A profile name must be provided.")
        self.profile_name = profile_name
        self.db = ModernJobDatabase()  # Consider dependency injection for better testing
        self.profile_dir = Path(f"profiles/{self.profile_name}")
        self.output_dir = self.profile_dir / "output"
        self.profile_dir.mkdir(parents=True, exist_ok=True)
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def customize_documents(self, job_data: Dict[str, Any]) -> Dict[str, str]:
        """
        Orchestrate the customization of resume and cover letter for a specific job.

        Args:
            job_data: A dictionary containing job details like 'company' and 'title'.

        Returns:
            A dictionary with paths to the generated documents.
        """
        try:
            company = job_data.get("company", "UnknownCompany")
            title = job_data.get("title", "UnknownPosition")
            
            resume_path = self._process_document("resume", company, title)
            cover_letter_path = self._process_document("cover_letter", company, title)

            return {
                "resume": str(resume_path),
                "cover_letter": str(cover_letter_path),
                "company": company,
                "job_title": title,
            }
        except Exception as e:
            logger.exception(f"Failed to customize documents for {company} - {title}: {e}")
            return {}

    def _process_document(self, doc_type: str, company: str, title: str) -> Optional[Path]:
        """
        Process a single document type (resume or cover letter).

        Handles finding the template, filling it, and converting to PDF.

        Args:
            doc_type: The type of document ('resume' or 'cover_letter').
            company: The target company name.
            title: The target job title.

        Returns:
            The path to the final generated document (PDF or DOCX).
        """
        sanitized_company = company.replace(' ', '_')
        sanitized_title = title.replace(' ', '_')
        
        template_path = self.profile_dir / f"{self.profile_name}_{doc_type.capitalize()}.docx"
        output_docx_path = self.output_dir / f"{doc_type}_{sanitized_company}_{sanitized_title}.docx"
        output_pdf_path = self.output_dir / f"{doc_type}_{sanitized_company}_{sanitized_title}.pdf"

        if template_path.exists():
            self._fill_docx_template(template_path, output_docx_path, company, title)
            if self._convert_docx_to_pdf(output_docx_path, output_pdf_path):
                return output_pdf_path
            return output_docx_path
        else:
            # Fallback to JSON/simple PDF generation if no .docx template
            json_content = self._create_custom_json_document(doc_type, company, title)
            self._write_pdf(output_pdf_path, json_content.get("content", ""), title=doc_type.capitalize())
            return output_pdf_path

    def _fill_docx_template(self, template_path: Path, output_path: Path, company: str, title: str):
        """
        Fill placeholders in a .docx template and save it to a new file.

        Args:
            template_path: Path to the .docx template.
            output_path: Path to save the filled .docx file.
            company: The company name to insert.
            title: The job title to insert.
        """
        try:
            doc = DocxDocument(template_path)
            for para in doc.paragraphs:
                # Use a loop to handle multiple replacements in the same paragraph run
                inline = para.runs
                for i in range(len(inline)):
                    text = inline[i].text
                    if "{company}" in text:
                        text = text.replace("{company}", company)
                    if "{job_title}" in text:
                        text = text.replace("{job_title}", title)
                    inline[i].text = text
            doc.save(output_path)
            logger.info(f"Successfully created DOCX: {output_path}")
        except Exception as e:
            logger.exception(f"Failed to fill DOCX template {template_path}: {e}")
            raise

    def _convert_docx_to_pdf(self, docx_path: Path, pdf_path: Path) -> bool:
        """
        Convert a .docx file to .pdf.

        Args:
            docx_path: Path to the source .docx file.
            pdf_path: Path to save the output .pdf file.

        Returns:
            True if conversion was successful, False otherwise.
        """
        try:
            docx2pdf_convert(str(docx_path), str(pdf_path))
            logger.info(f"Successfully converted {docx_path.name} to PDF.")
            return True
        except Exception as e:
            logger.warning(f"docx2pdf conversion failed for {docx_path.name}: {e}. Fallback may be needed.")
            return False

    def _create_custom_json_document(self, doc_type: str, company: str, title: str) -> Dict[str, Any]:
        """
        Create a customized document as a JSON object.

        Args:
            doc_type: The type of document ('resume' or 'cover_letter').
            company: The target company name.
            title: The target job title.

        Returns:
            A dictionary representing the customized document.
        """
        content = {
            "profile_name": self.profile_name,
            "target_company": company,
            "target_position": title,
            "customization_date": datetime.now().isoformat(),
            "content": f"Customized {doc_type.replace('_', ' ')} for {title} at {company}",
        }
        output_file = self.output_dir / f"customized_{doc_type}.json"
        try:
            with output_file.open("w", encoding="utf-8") as f:
                json.dump(content, f, indent=2)
            logger.info(f"✅ Customized {doc_type} JSON created: {output_file}")
        except IOError as e:
            logger.exception(f"Error writing {doc_type} JSON file: {e}")
        return content

    def _write_pdf(self, pdf_path: Path, content: str, title: str = "Document"):
        """
        Write content to a simple PDF file with basic formatting.

        Args:
            pdf_path: The path to save the PDF file.
            content: The text content to write to the PDF.
            title: The title of the document.
        """
        try:
            c = canvas.Canvas(str(pdf_path), pagesize=LETTER)
            width, height = LETTER
            margin = 72  # 1 inch

            c.setFont("Helvetica-Bold", 16)
            c.drawString(margin, height - margin, title)

            c.setFont("Helvetica", 12)
            text = c.beginText(margin, height - margin - 30)
            text.setFont("Helvetica", 12)
            
            for line in content.split('\n'):
                text.textLine(line)

            c.drawText(text)
            c.save()
            logger.info(f"Successfully wrote PDF: {pdf_path}")
        except Exception as e:
            logger.exception(f"Failed to write PDF {pdf_path}: {e}")
            raise

    # --- AI-Powered Document Generation Methods ---

    def get_available_templates(self) -> List[str]:
        """Return a list of available document templates."""
        try:
            templates = []
            
            # Check for DOCX templates in profile directory
            docx_templates = list(self.profile_dir.glob("*.docx"))
            for template in docx_templates:
                templates.append(template.name)
            
            # Always include AI-generated templates
            templates.extend([
                "AI-Generated Resume (Gemini)",
                "AI-Generated Cover Letter (Gemini)",
                "Default PDF Template"
            ])
            
            logger.info(f"Found {len(templates)} available templates")
            return templates
            
        except Exception as e:
            logger.error(f"Error discovering templates: {e}")
            return ["Default PDF Template"]

    def generate_ai_cover_letter(self, job_data: Dict[str, Any], profile_data: Dict[str, Any]) -> str:
        """Generate an AI-powered cover letter using Gemini API."""
        try:
            from src.utils.gemini_client import GeminiClient
            from src.utils.pdf_generator import PDFGenerator
            
            # Initialize Gemini client
            gemini = GeminiClient()
            
            # Generate cover letter content
            logger.info(f"Generating AI cover letter for {job_data.get('company', 'Unknown')} - {job_data.get('title', 'Unknown')}")
            cover_letter_content = gemini.generate_cover_letter(profile_data, job_data)
            
            if not cover_letter_content:
                raise ValueError("Gemini API returned empty content")
            
            # Save as text file first
            company = job_data.get('company', 'UnknownCompany').replace(' ', '_')
            title = job_data.get('title', 'UnknownPosition').replace(' ', '_')
            
            text_path = self.output_dir / f"cover_letter_{company}_{title}.txt"
            pdf_path = self.output_dir / f"cover_letter_{company}_{title}.pdf"
            
            # Save text content
            with open(text_path, 'w', encoding='utf-8') as f:
                f.write(cover_letter_content)
            
            # Generate PDF
            pdf_generator = PDFGenerator()
            success = pdf_generator.text_to_pdf(cover_letter_content, pdf_path, "cover_letter")
            
            if success:
                logger.info(f"✅ AI cover letter generated successfully: {pdf_path}")
                return str(pdf_path)
            else:
                logger.warning(f"PDF generation failed, returning text file: {text_path}")
                return str(text_path)
                
        except Exception as e:
            logger.error(f"Failed to generate AI cover letter: {e}")
            # Fallback to simple template
            return self._generate_fallback_cover_letter(job_data, profile_data)

    def generate_ai_resume(self, job_data: Dict[str, Any], profile_data: Dict[str, Any]) -> str:
        """Generate an AI-powered resume using Gemini API."""
        try:
            from src.utils.gemini_client import GeminiClient
            from src.utils.pdf_generator import PDFGenerator
            
            # Initialize Gemini client
            gemini = GeminiClient()
            
            # Generate resume content
            logger.info(f"Generating AI resume for {job_data.get('company', 'Unknown')} - {job_data.get('title', 'Unknown')}")
            resume_content = gemini.generate_resume(profile_data, job_data)
            
            if not resume_content:
                raise ValueError("Gemini API returned empty content")
            
            # Save as text file first
            company = job_data.get('company', 'UnknownCompany').replace(' ', '_')
            title = job_data.get('title', 'UnknownPosition').replace(' ', '_')
            
            text_path = self.output_dir / f"resume_{company}_{title}.txt"
            pdf_path = self.output_dir / f"resume_{company}_{title}.pdf"
            
            # Save text content
            with open(text_path, 'w', encoding='utf-8') as f:
                f.write(resume_content)
            
            # Generate PDF
            pdf_generator = PDFGenerator()
            success = pdf_generator.text_to_pdf(resume_content, pdf_path, "resume")
            
            if success:
                logger.info(f"✅ AI resume generated successfully: {pdf_path}")
                return str(pdf_path)
            else:
                logger.warning(f"PDF generation failed, returning text file: {text_path}")
                return str(text_path)
                
        except Exception as e:
            logger.error(f"Failed to generate AI resume: {e}")
            # Fallback to simple template
            return self._generate_fallback_resume(job_data, profile_data)

    def _generate_fallback_cover_letter(self, job_data: Dict[str, Any], profile_data: Dict[str, Any]) -> str:
        """Generate a simple fallback cover letter when AI generation fails."""
        try:
            company = job_data.get('company', 'Hiring Manager')
            title = job_data.get('title', 'the position')
            name = profile_data.get('name', 'Nirajan Khadka')
            
            content = f"""Dear {company} Hiring Team,

I am writing to express my strong interest in the {title} position at {company}. With my background in data analysis and technical skills, I am confident that I would be a valuable addition to your team.

My experience includes:
- Data analysis using Python, SQL, and various BI tools
- Creating automated dashboards and reports
- Collaborating with cross-functional teams
- Developing predictive models and insights

I am excited about the opportunity to contribute to {company}'s continued success and would welcome the chance to discuss how my skills and experience align with your needs.

Thank you for your consideration.

Sincerely,
{name}"""
            
            company_safe = company.replace(' ', '_')
            title_safe = title.replace(' ', '_')
            fallback_path = self.output_dir / f"cover_letter_fallback_{company_safe}_{title_safe}.txt"
            
            with open(fallback_path, 'w', encoding='utf-8') as f:
                f.write(content)
            
            logger.info(f"Generated fallback cover letter: {fallback_path}")
            return str(fallback_path)
            
        except Exception as e:
            logger.error(f"Failed to generate fallback cover letter: {e}")
            return ""

    def _generate_fallback_resume(self, job_data: Dict[str, Any], profile_data: Dict[str, Any]) -> str:
        """Generate a simple fallback resume when AI generation fails."""
        try:
            name = profile_data.get('name', 'Nirajan Khadka')
            email = profile_data.get('email', 'your.email@example.com')
            phone = profile_data.get('phone', '(555) 123-4567')
            location = profile_data.get('location', 'Your City, State')
            
            content = f"""{name}
{email} | {phone} | {location}

PROFESSIONAL SUMMARY
{profile_data.get('summary', 'Experienced professional with strong analytical and technical skills.')}

EXPERIENCE
{profile_data.get('experience', 'Please see attached resume for detailed experience.')}

SKILLS
{profile_data.get('skills', 'Python, SQL, Data Analysis, Business Intelligence')}

EDUCATION
{profile_data.get('education', 'Bachelor of Science in relevant field')}

CERTIFICATIONS
{profile_data.get('certifications', 'Industry-relevant certifications')}"""
            
            company = job_data.get('company', 'UnknownCompany').replace(' ', '_')
            title = job_data.get('title', 'UnknownPosition').replace(' ', '_')
            fallback_path = self.output_dir / f"resume_fallback_{company}_{title}.txt"
            
            with open(fallback_path, 'w', encoding='utf-8') as f:
                f.write(content)
            
            logger.info(f"Generated fallback resume: {fallback_path}")
            return str(fallback_path)
            
        except Exception as e:
            logger.error(f"Failed to generate fallback resume: {e}")
            return ""


def customize(job_data: Dict[str, Any], profile_name: str) -> Dict[str, str]:
    """
    Convenience function to customize documents for a job.

    Args:
        job_data: A dictionary containing job details.
        profile_name: The name of the profile to use.

    Returns:
        A dictionary with paths to the customized documents.
    """
    modifier = DocumentModifier(profile_name)
    return modifier.customize_documents(job_data)
